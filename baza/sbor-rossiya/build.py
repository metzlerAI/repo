import json, re, datetime, collections, sys
from prof import ok
from urllib.parse import unquote
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
TODAY = datetime.date(2026, 10, 17)
LIVE = TODAY - datetime.timedelta(days=183)   # фильтр: последний отзыв не старше 6 месяцев
FRESH = TODAY - datetime.timedelta(days=int(sys.argv[1]) if len(sys.argv) > 1 else 90)  # для A
DAYS = [('Mon','пн'),('Tue','вт'),('Wed','ср'),('Thu','чт'),('Fri','пт'),('Sat','сб'),('Sun','вс')]
def hours(s):
    if not s: return ''
    sl = ['; '.join(f"{h['from']}-{h['to']}" for h in (s.get(d) or {}).get('working_hours', [])) or 'выходной' for d,_ in DAYS]
    if len(set(sl)) == 1: return 'круглосуточно' if sl[0] == '00:00-24:00' else 'ежедневно ' + sl[0]
    out, i = [], 0
    while i < 7:
        j = i
        while j + 1 < 7 and sl[j+1] == sl[i]: j += 1
        out.append(DAYS[i][1] + ('-' + DAYS[j][1] if j > i else '') + ' ' + sl[i]); i = j + 1
    return ', '.join(out)
c = json.load(open('cards.json'))
dup = set(); seen = len(c); rows = []; why = collections.Counter()
for id, x in c.items():
    if not ok(x['n'], x['r']): why['не профиль'] += 1; continue
    key = (x['n'], x['city'])
    if key in dup: why['дубль'] += 1; continue
    dup.add(key)
    if 'err' in x: why['ошибка'] += 1; continue
    if x['s'] == 'есть' or 'clck.ru' in x.get('site', ''): why['есть сайт'] += 1; continue
    if not (x['tg'] or x['wa']): why['нет мессенджера'] += 1; continue
    dates = [r['date'] for r in x['reviews'] if r['date']]
    last = max(dates) if dates else ''
    if last and datetime.date.fromisoformat(last) < LIVE: why['отзывы старые'] += 1; continue
    own = [p for p in x['photos'] if p['src'] in ('2gis_owner', 'instagram')]
    isA = x['rv'] >= 15 and last and datetime.date.fromisoformat(last) >= FRESH and len(own) >= 3
    site = x['site'].lower()
    sk = 'нет' if x['s'] == 'нет' else ('taplink' if 'taplink' in site else 'соцсети') + ' (' + x['site'] + ')'
    wa = re.sub(r'\?.*', '', x['wa']); tg = x['tg']
    rows.append(dict(pr='A' if isA else 'B', city=x['city'], n=x['n'], r=x['r'], rv=x['rv'], rt=x['rt'], last=last or 'не видно', site=sk,
        tg=tg, wa=wa, ph=x['ph'], adr=x['adr'], url=f"https://2gis.ru/{x['slug']}/firm/{id}", own=own, revs=x['reviews'], hrs=hours(x['sched'])))
CHK = json.load(open('checks.json'))
for r in rows: r['chk'] = CHK.get(r['url'].rsplit('/', 1)[1], {}).get('note', 'не проверялось')
cut = [r for r in rows if CHK.get(r['url'].rsplit('/', 1)[1], {}).get('ex')]
rows = [r for r in rows if r not in cut]
rows.sort(key=lambda r: (r['pr'], -r['rv']))
A = [r for r in rows if r['pr'] == 'A']; B = [r for r in rows if r['pr'] == 'B']
print('просмотрено', seen, 'прошли', len(rows), 'A', len(A), 'B', len(B), dict(why))
print('города', collections.Counter(r['city'] for r in rows).most_common(8))
if '--write' not in sys.argv: sys.exit()
main = (A + B)[:50]; rest = (A + B)[50:]
def pick_reviews(rv):
    good = [r for r in rv if r['rating'] == 5 and r['text'] and 40 <= len(r['text']) <= 260 and r['name']]
    return '\n\n'.join(f"«{r['text'].strip()}» ({r['name']}, {r['date']})" for r in good[:3])
H = ['приоритет','город','название','рубрика','отзывов','рейтинг','дата последнего отзыва','сайт (нет / соцсети / taplink)','Telegram','WhatsApp','телефон','адрес','ссылка на карточку 2ГИС','фото (ссылки)','отзывы','часы работы','проверка сайта']
wb = Workbook()
def sheet(ws, L, extra):
    ws.append(H)
    for k, r in enumerate(L):
        full = extra and r['pr'] == 'A' and k < 15
        ws.append([r['pr'], r['city'], r['n'], r['r'], r['rv'], r['rt'], r['last'], r['site'], r['tg'], r['wa'], r['ph'], r['adr'], r['url'],
                   '\n'.join(p['url'] for p in r['own'][:6]) if full else '', pick_reviews(r['revs']) if full else '', r['hrs'] if full else '', r['chk']])
    for cell in ws[1]: cell.font = Font(bold=True)
    for col, w in zip('ABCDEFGHIJKLMNOPQ', [6,14,28,22,8,8,12,24,22,16,18,30,40,60,80,30,40]): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
ws = wb.active; ws.title = 'база 50'; sheet(ws, main, True)
ws2 = wb.create_sheet('резерв'); sheet(ws2, rest, False)
ws3 = wb.create_sheet('есть сайт (убраны)'); sheet(ws3, cut, False)
wb.save('baza-rossiya.xlsx')
print('в файле', len(main), 'A', sum(r['pr']=='A' for r in main), 'резерв', len(rest))
print('A15 без 3 отзывов:', [r['n'] for r in main[:15] if pick_reviews(r['revs']).count('«') < 2], 'без часов:', [r['n'] for r in main[:15] if not r['hrs']])
